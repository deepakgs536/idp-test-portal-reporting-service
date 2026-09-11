from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font


def export_candidates_to_excel(candidates):
    """
    Generate an Excel file from candidate reports.

    Args:
        candidates (list): List of candidate report dictionaries.

    Returns:
        bytes: Excel file as bytes.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Candidate Reports"

    headers = [
        "Candidate Name",
        "Email",
        "College",
        "Mobile",
        "Score",
        "Total Marks",
        "Percentage",
        "Section Scores",
        "Correct Answers",
        "Wrong Answers",
        "Unanswered",
        "Time Taken (sec)",
        "Status",
        "Warning Count",
        "Submitted At",
        "Generated At"
    ]

    # Header Row
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)

    # Data Rows
    for row, report in enumerate(candidates, start=2):
        proctor = report.get("proctoringDetails", {})

        sheet.cell(row=row, column=1).value = report.get("candidateName")
        sheet.cell(row=row, column=2).value = report.get("mailId")
        sheet.cell(row=row, column=3).value = report.get("college")
        sheet.cell(row=row, column=4).value = report.get("mobile")
        sheet.cell(row=row, column=5).value = float(report.get("score", 0))
        sheet.cell(row=row, column=6).value = float(report.get("totalMarks", 0))
        sheet.cell(row=row, column=7).value = float(report.get("percentage", 0))
        
        section_scores = []
        for sec_name, sec_score in report.get("sectionScores", {}).items():
            section_scores.append(f"{sec_name}: {float(sec_score)}")
            
        sheet.cell(row=row, column=8).value = " | ".join(section_scores)
        sheet.cell(row=row, column=9).value = report.get("correctAnswers")
        sheet.cell(row=row, column=10).value = report.get("wrongAnswers")
        sheet.cell(row=row, column=11).value = report.get("unanswered")
        sheet.cell(row=row, column=12).value = float(report.get("timeTaken", 0))
        sheet.cell(row=row, column=13).value = report.get("status")
        sheet.cell(row=row, column=14).value = proctor.get("warningCount", 0)
        sheet.cell(row=row, column=15).value = report.get("submittedAt")
        sheet.cell(row=row, column=16).value = report.get("generatedAt")

    # Auto-fit column width
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 3, 40)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue()